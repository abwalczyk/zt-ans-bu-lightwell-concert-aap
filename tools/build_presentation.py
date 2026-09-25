#!/usr/bin/env python3
"""Build the lab intro deck as a .pptx styled to the Red Hat Summit 2026 template.

Google Slides imports .pptx directly, and because every diagram here is built
from native shapes (not a rendered image) it stays editable after the import.

Style is taken from the Summit 2026 / AnsibleFest lab template:
  * dark Summit title background (extracted asset), white display type
  * white content slides with the left brand bracket, centred near-black title
    and a centred red subtitle
  * numbered steps as light grey circles with a grey ring
  * step cards as white boxes with a thin red rounded border
  * page number bottom-left, Red Hat logo bottom-right

Typeface is Red Hat Display (titles) / Red Hat Text (body). Both ship as Google
Fonts, so they resolve after importing into Google Slides.

Usage:
    pip install python-pptx
    python3 tools/build_presentation.py [output.pptx]

Source of truth for the wording is PRESENTATION.md -- keep the two in sync.
"""

import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ASSETS = os.path.join(HERE, "assets")
TITLE_BG = os.path.join(ASSETS, "summit-title-bg.png")
RH_LOGO = os.path.join(ASSETS, "redhat-logo.png")

# ── Summit palette ─────────────────────────────────────────────────────
RED = RGBColor(0xEE, 0x00, 0x00)        # Red Hat red
BRACKET = RGBColor(0xA3, 0x00, 0x00)    # left brand bracket
INK = RGBColor(0x15, 0x15, 0x15)        # near-black body/title
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
FAINT = RGBColor(0x8C, 0x8C, 0x8C)
HAIRLINE = RGBColor(0xBB, 0xBB, 0xBB)
CIRCLE_FILL = RGBColor(0xF2, 0xF2, 0xF2)
CIRCLE_RING = RGBColor(0x9C, 0x9C, 0x9C)
PINK = RGBColor(0xF9, 0xD5, 0xD8)       # highlight behind key phrases
PLUM = RGBColor(0x21, 0x13, 0x4D)       # Summit dark purple
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
WASH = RGBColor(0xFA, 0xFA, 0xFA)

DISPLAY = "Red Hat Display"
TEXT = "Red Hat Text"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


# ── Primitives ─────────────────────────────────────────────────────────
def textbox(slide, x, y, w, h, text, size=12, bold=False, color=INK,
            align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False,
            spacing=1.0, font=TEXT):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = anchor
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        run = p.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        run.font.name = font
    return box


def _no_shadow(shp):
    shp.shadow.inherit = False
    return shp


def card(slide, x, y, w, h, outline=RED, fill=WHITE, dash=None, width=Pt(1.25),
         radius=0.06):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.adjustments[0] = radius
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if outline is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = outline
        shp.line.width = width
        if dash:
            shp.line.dash_style = dash
    _no_shadow(shp)
    shp.text_frame.text = ""
    return shp


def step_card(slide, x, y, w, h, title, body="", outline=RED, title_size=10,
              body_size=8.5, title_color=INK, dash=None):
    """White card, thin red rounded border -- the Summit diagram unit."""
    shp = card(slide, x, y, w, h, outline=outline, dash=dash)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.07)
    tf.margin_top = tf.margin_bottom = Inches(0.05)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE

    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.line_spacing = 0.95
    r = p.add_run()
    r.text = title
    r.font.size = Pt(title_size)
    r.font.bold = True
    r.font.color.rgb = title_color
    r.font.name = TEXT

    if body:
        p2 = tf.add_paragraph()
        p2.alignment = PP_ALIGN.CENTER
        p2.line_spacing = 0.95
        p2.space_before = Pt(3)
        r2 = p2.add_run()
        r2.text = body
        r2.font.size = Pt(body_size)
        r2.font.color.rgb = MUTED
        r2.font.name = TEXT
    return shp


def numbered_circle(slide, cx, cy, n, d=Inches(0.34)):
    """Summit style: light grey disc, grey ring, black numeral."""
    shp = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, int(cx - d / 2), int(cy - d / 2), d, d)
    shp.fill.solid()
    shp.fill.fore_color.rgb = CIRCLE_FILL
    shp.line.color.rgb = CIRCLE_RING
    shp.line.width = Pt(1)
    _no_shadow(shp)
    tf = shp.text_frame
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = str(n)
    r.font.size = Pt(12)
    r.font.bold = True
    r.font.color.rgb = INK
    r.font.name = DISPLAY
    return shp


def arrow(slide, x1, y1, x2, y2, color=RGBColor(0x4D, 0x4D, 0x4D),
          width=Pt(1.75)):
    conn = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, int(x1), int(y1), int(x2), int(y2))
    conn.line.color.rgb = color
    conn.line.width = width
    ln = conn.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd"))
    tail.set("type", "triangle")
    tail.set("w", "med")
    tail.set("len", "med")
    return conn


def highlight_run(run, color=PINK):
    """Pink marker behind a run -- the template's emphasis treatment."""
    rPr = run._r.get_or_add_rPr()
    hl = etree.SubElement(rPr, qn("a:highlight"))
    clr = etree.SubElement(hl, qn("a:srgbClr"))
    clr.set("val", "%02X%02X%02X" % (color[0], color[1], color[2]))
    return run


# ── Slide chrome ───────────────────────────────────────────────────────
def brand_bracket(slide):
    """The broken vertical rule at the left edge, top and bottom segments."""
    for top, height in ((Inches(0.0), Inches(1.0)),
                        (Inches(7.14), Inches(0.36))):
        bar = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0.50), top, Pt(1.5), height)
        bar.fill.solid()
        bar.fill.fore_color.rgb = BRACKET
        bar.line.fill.background()
        _no_shadow(bar)
        bar.text_frame.text = ""


def chrome(slide, page=None):
    brand_bracket(slide)
    if os.path.exists(RH_LOGO):
        slide.shapes.add_picture(
            RH_LOGO, Inches(11.71), Inches(6.90), Inches(1.08))
    if page is not None:
        textbox(slide, Inches(0.42), Inches(6.80), Inches(0.5), Inches(0.2),
                str(page), size=8, color=FAINT)


def content_slide(prs, title, subtitle=None, page=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    chrome(s, page)
    textbox(s, Inches(1.1), Inches(0.42), Inches(11.13), Inches(0.5), title,
            size=27, color=INK, align=PP_ALIGN.CENTER, font=DISPLAY,
            spacing=0.95)
    if subtitle:
        textbox(s, Inches(1.1), Inches(0.97), Inches(11.13), Inches(0.35),
                subtitle, size=16, color=RED, align=PP_ALIGN.CENTER,
                font=TEXT, spacing=0.95)
    return s


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


# ── Slides ─────────────────────────────────────────────────────────────
def cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if os.path.exists(TITLE_BG):
        s.shapes.add_picture(TITLE_BG, 0, 0, SLIDE_W, SLIDE_H)
    else:
        bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        bg.fill.solid()
        bg.fill.fore_color.rgb = PLUM
        bg.line.fill.background()
        _no_shadow(bg)
        bg.text_frame.text = ""

    textbox(s, Inches(0.78), Inches(2.55), Inches(9.6), Inches(1.6),
            "Lab: Proactive dependency\nremediation", size=40, color=WHITE,
            font=DISPLAY, spacing=1.05)
    textbox(s, Inches(0.80), Inches(4.30), Inches(9.6), Inches(0.4),
            "Start from the fix, not the flaw.", size=19, color=WHITE,
            italic=True, font=TEXT)
    textbox(s, Inches(0.80), Inches(5.05), Inches(10.2), Inches(0.9),
            "Red Hat Lightwell  ·  IBM Concert  ·  Ansible Automation "
            "Platform  ·  Tekton  ·  ArgoCD\n~10 minute intro  ·  "
            "55 minutes hands-on",
            size=13, color=WHITE, font=TEXT, spacing=1.5)
    notes(s, "Set the frame: this lab is not about finding vulnerabilities. "
             "It is about what happens the instant a trusted fix exists.")
    return s


def slide_problem(prs):
    s = content_slide(prs, "Patching is reactive, and slow",
                      "The bottleneck is not the fix — it is knowing where "
                      "the fix needs to go", page=2)

    items = [
        ("A CVE drops. Now what?", True),
        ("Which of our applications actually use the affected package?", False),
        ("Who owns them? What breaks downstream if we patch?", False),
        ("Someone opens a spreadsheet. Weeks pass.", False),
    ]
    y = Inches(1.85)
    for text, is_head in items:
        if is_head:
            textbox(s, Inches(1.15), y, Inches(6.6), Inches(0.35), text,
                    size=17, bold=True, font=DISPLAY)
            y += Inches(0.55)
        else:
            textbox(s, Inches(1.20), y, Inches(0.3), Inches(0.3), "›",
                    size=15, bold=True, color=RED)
            textbox(s, Inches(1.55), y, Inches(6.2), Inches(0.6), text,
                    size=14, spacing=1.1)
            y += Inches(0.52)

    y += Inches(0.25)
    textbox(s, Inches(1.15), y, Inches(6.7), Inches(0.9),
            "Industry average to remediate a critical dependency CVE is "
            "measured in weeks, not minutes.", size=14, spacing=1.15)

    # Right: the manual loop
    card(s, Inches(8.45), Inches(1.85), Inches(3.9), Inches(3.9),
         outline=HAIRLINE, fill=WASH)
    textbox(s, Inches(8.78), Inches(2.10), Inches(3.3), Inches(0.3),
            "THE MANUAL LOOP", size=11, bold=True, color=RED, font=DISPLAY)
    steps = ["Scanner flags a CVE", "Hunt for affected apps",
             "Chase down owners", "Open tickets", "Manual PR + deploy",
             "Hope nothing broke"]
    yy = Inches(2.58)
    for i, stp in enumerate(steps):
        numbered_circle(s, Inches(9.03), yy + Inches(0.13), i + 1,
                        d=Inches(0.26))
        textbox(s, Inches(9.30), yy, Inches(2.9), Inches(0.3), stp, size=12.5)
        yy += Inches(0.43)
    textbox(s, Inches(8.78), Inches(5.28), Inches(3.4), Inches(0.3),
            "Elapsed: days to weeks", size=13, bold=True, color=RED,
            font=DISPLAY)

    notes(s, "Most vulnerability programs are a detection pipeline bolted "
             "onto a manual remediation process. We're going to invert that.")
    return s


def slide_inversion(prs):
    s = content_slide(prs, "Start from the fix, not the flaw",
                      "Move the trigger, and everything downstream changes",
                      page=3)

    heads = ["Reactive (typical)", "Proactive (this lab)"]
    rows = [
        ("Trigger: a scanner finds a CVE",
         "Trigger: Lightwell publishes a remediated package"),
        ("“We have a problem”", "“We have a solution — go apply it”"),
        ("Hunt for affected apps", "Concert already knows"),
        ("Manual PR, manual deploy", "Tekton builds, ArgoCD deploys"),
    ]

    colx = [Inches(1.15), Inches(7.0)]
    colw = Inches(5.2)
    for i, head in enumerate(heads):
        bar = card(s, colx[i], Inches(1.72), colw, Inches(0.44),
                   outline=None, fill=RED if i else RGBColor(0x6A, 0x6A, 0x6A))
        tf = bar.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = head
        r.font.size = Pt(13.5)
        r.font.bold = True
        r.font.color.rgb = WHITE
        r.font.name = DISPLAY

    y = Inches(2.34)
    for a, b in rows:
        for i, val in enumerate((a, b)):
            c = card(s, colx[i], y, colw, Inches(0.66),
                     outline=HAIRLINE if i == 0 else RED, fill=WHITE)
            tf = c.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_right = Inches(0.14)
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            r = p.add_run()
            r.text = val
            r.font.size = Pt(12.5)
            r.font.color.rgb = INK
            r.font.name = TEXT
        arrow(s, colx[0] + colw + Inches(0.08), y + Inches(0.33),
              colx[1] - Inches(0.08), y + Inches(0.33), width=Pt(1.5))
        y += Inches(0.78)

    textbox(s, Inches(1.15), Inches(5.72), Inches(11.05), Inches(0.9),
            "Red Hat Lightwell publishes a hardened, remediated build of a "
            "package into Artifactory. That publish event is the trigger — "
            "the fix arrives before anyone files a ticket.",
            size=14, spacing=1.2, align=PP_ALIGN.CENTER)

    notes(s, "This is the key idea of the lab. Everything downstream follows "
             "from moving the trigger.")
    return s


def slide_cast(prs):
    s = content_slide(prs, "The cast: who does what",
                      "Concert answers “who is affected?” — AO owns the "
                      "decisions", page=4)

    rows = [
        ("Artifactory", "Package store — emits the webhook when a Lightwell "
                        "package lands"),
        ("Event-Driven Ansible", "Always-on listener — turns the event into a "
                                 "workflow"),
        ("IBM Concert", "Pre-indexed SBOM inventory + Arena View dependency "
                        "topology"),
        ("Automation Orchestrator", "The brain — routing decisions, approval "
                                    "gates, audit trail"),
        ("AAP", "The hands — playbook execution"),
        ("Tekton", "CI — build, test, image, GitOps manifest bump"),
        ("ArgoCD", "CD — GitOps sync to OpenShift, drift detection, rollback"),
    ]

    y = Inches(1.72)
    for name, role in rows:
        chev = slide_chevron(s, Inches(1.15), y, Inches(3.5), Inches(0.58),
                             name)
        textbox(s, Inches(5.05), y + Inches(0.13), Inches(7.2), Inches(0.4),
                role, size=13, color=INK)
        y += Inches(0.68)

    textbox(s, Inches(1.15), Inches(6.62), Inches(11.05), Inches(0.35),
            "Note what Concert is not doing: it is not orchestrating.",
            size=12.5, italic=True, color=MUTED, align=PP_ALIGN.CENTER)

    notes(s, "Concert is an inventory and topology service here, not the "
             "orchestrator. AO owns routing and policy.")
    return s


def slide_chevron(slide, x, y, w, h, label, fill=WHITE, outline=HAIRLINE,
                  color=INK, size=13):
    """Right-pointing chevron card, as used on the template's takeaway slide."""
    shp = slide.shapes.add_shape(MSO_SHAPE.PENTAGON, x, y, w, h)
    shp.adjustments[0] = 0.18
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = outline
    shp.line.width = Pt(1.25)
    _no_shadow(shp)
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.16)
    tf.margin_right = Inches(0.3)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = label
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = TEXT
    return shp


def slide_architecture(prs):
    """The swimlane. Ansible lane on top, IBM Concert lane beneath."""
    s = content_slide(prs, "Lab architecture",
                      "From package publish to validated production deploy",
                      page=5)

    left = Inches(0.78)
    lane_x = Inches(1.80)
    col_w = Inches(1.42)
    gap = Inches(0.21)
    ncols = 7
    lane_w = ncols * col_w + (ncols - 1) * gap

    top_y, bot_y = Inches(2.62), Inches(4.52)
    box_h = Inches(1.14)
    lane_pad = Inches(0.20)
    band_w = int(lane_w + (lane_x - left))

    def col_x(i):
        return int(lane_x + i * (col_w + gap))

    def lane(y, label, accent):
        rect = card(s, left, int(y - lane_pad), band_w,
                    int(box_h + 2 * lane_pad), outline=HAIRLINE, fill=WASH,
                    dash=MSO_LINE_DASH_STYLE.ROUND_DOT, width=Pt(1),
                    radius=0.04)
        tab = card(s, left, int(y - lane_pad), Inches(0.07),
                   int(box_h + 2 * lane_pad), outline=None, fill=accent,
                   radius=0.2)
        textbox(s, left + Inches(0.16), int(y - lane_pad) + Inches(0.06),
                Inches(0.84), int(box_h + 2 * lane_pad) - Inches(0.12), label,
                size=9.5, bold=True, color=accent, anchor=MSO_ANCHOR.MIDDLE,
                spacing=0.9, font=DISPLAY)
        return rect

    lane(top_y, "Ansible\nAutomation\nPlatform", RED)
    lane(bot_y, "IBM\nConcert", PLUM)

    headers = ["Package publish", "Event detection", "Orchestration",
               "Impact analysis", "Source update", "Build & deploy",
               "Validation"]
    for i, h in enumerate(headers):
        textbox(s, col_x(i), Inches(1.58), col_w, Inches(0.5), h, size=10.5,
                bold=True, align=PP_ALIGN.CENTER, spacing=0.9, font=TEXT)
        numbered_circle(s, col_x(i) + col_w / 2, Inches(2.18), i + 1)

    # Step 1 spans both lanes: the package publish is the single event that
    # both the Ansible side and the Concert side react to.
    span_h = int(bot_y + box_h - top_y)
    trigger_box = step_card(s, col_x(0), int(top_y), col_w, span_h,
                            "Artifactory",
                            "lightwell-remediated\nrepo emits webhook\n\n"
                            "One publish event,\nboth lanes react",
                            outline=RED)

    # Solid = the path this lab walks. Dotted = the same step, owned by
    # Concert (or by AO) instead -- the alternatives we want to talk about.
    cells = [
        # (col, lane, alt?, title, body)
        (1, "t", False, "Event-Driven\nAnsible", "Rulebook matches,\nfires workflow"),
        (2, "t", False, "Automation\nOrchestrator", "Owns routing,\npolicy, audit"),
        (3, "t", True, "Automation\nOrchestrator", "Could own the query\nand the policy gate"),
        (4, "t", False, "AAP playbook", "Bump requirements,\nbranch, commit, push"),
        (5, "t", False, "Tekton + ArgoCD", "Test, build image,\nGitOps sync"),
        (6, "t", False, "AAP validation", "3/3 pods healthy,\n/health reports 6.0.2"),

        (1, "b", True, "Concert ingest", "New Lightwell build\nindexed in the SBOM"),
        (2, "b", True, "Risk prioritization", "Ranks affected apps\nby exposure"),
        (3, "b", False, "SBOM inventory\n+ Arena View", "12 of 537 apps\non an older build"),
        (4, "b", True, "Concert\nSecure Coder", "Safe-version fix,\nopens a pull request"),
        (5, "b", True, "Concert Workflows", "Change request,\napproval, tracked"),
        (6, "b", False, "SBOM updated", "Portfolio drift\n12 to 11"),
    ]

    grid = {}
    for i, ln, alt, title, body in cells:
        y = top_y if ln == "t" else bot_y
        grid[(i, ln)] = step_card(
            s, col_x(i), int(y), col_w, box_h, title, body,
            outline=RED if ln == "t" else PLUM,
            dash=MSO_LINE_DASH_STYLE.ROUND_DOT if alt else None,
            title_color=MUTED if alt else INK)

    def edge(b, side):
        if side == "r":
            return b.left + b.width, b.top + b.height // 2
        return b.left, b.top + b.height // 2

    # Main path: straight along the Ansible lane, step 1 through 7.
    chain = [trigger_box] + [grid[(i, "t")] for i in range(1, 7)]
    for a, b in zip(chain, chain[1:]):
        _, y = edge(grid[(1, "t")], "r")
        arrow(s, a.left + a.width + Inches(0.02), y, b.left - Inches(0.02), y)

    # Artifactory also feeds the Concert lane directly.
    _, by = edge(grid[(1, "b")], "l")
    arrow(s, trigger_box.left + trigger_box.width + Inches(0.02), by,
          grid[(1, "b")].left - Inches(0.02), by)

    # Column 4: AO asks Concert who is affected, Concert answers.
    q = grid[(3, "t")]
    a4 = grid[(3, "b")]
    down_x = q.left + q.width // 2 - Inches(0.26)
    up_x = q.left + q.width // 2 + Inches(0.26)
    arrow(s, down_x, q.top + q.height + Inches(0.02),
          down_x, a4.top - Inches(0.02))
    arrow(s, up_x, a4.top - Inches(0.02), up_x, q.top + q.height + Inches(0.02))

    # Column 7: validation writes back down into Concert's inventory.
    v = grid[(6, "t")]
    w = grid[(6, "b")]
    arrow(s, v.left + v.width // 2, v.top + v.height + Inches(0.02),
          w.left + w.width // 2, w.top - Inches(0.02))

    # TRIGGER ribbon on step 1
    trig = card(s, col_x(0) + Inches(0.29), int(top_y - Inches(0.13)),
                Inches(0.84), Inches(0.25), outline=None, fill=RED,
                radius=0.30)
    tf = trig.text_frame
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = "TRIGGER"
    r.font.size = Pt(7.5)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = DISPLAY

    textbox(s, left, Inches(5.94), band_w, Inches(0.26),
            "Solid outline — the path this lab walks   ·   "
            "Dotted outline — the same step, owned by Concert instead",
            size=9.5, color=FAINT, italic=True, align=PP_ALIGN.CENTER)

    foot = textbox(s, left, Inches(6.28), band_w, Inches(0.4), "", size=13.5,
                   align=PP_ALIGN.CENTER)
    p = foot.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    for txt, hl, bold in [
            ("The trigger is a fix being published", True, True),
            (" — not a flaw being found.  Publish to validated production "
             "deploy in ~3 minutes.", False, False)]:
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(13.5)
        r.font.bold = bold
        r.font.color.rgb = INK
        r.font.name = TEXT
        if hl:
            highlight_run(r)

    notes(s,
          "Walk it left to right. The thing to notice is step 1: the trigger "
          "is a fix being published, not a flaw being found. Concert owns the "
          "bottom lane -- it answers 'who is affected' in one API call and "
          "then gets written back to at the end so inventory reflects "
          "reality. Everything in the top lane is Ansible orchestrating and "
          "executing, with Tekton and ArgoCD doing the actual build and "
          "deploy. The dotted boxes are the conversation piece: Concert can "
          "own most of these steps itself -- Secure Coder opens the fix PR, "
          "Concert Workflows drive the change request and approval. We run "
          "them through AAP here because that is where the policy and the "
          "audit trail live, but the split is a design choice, not a "
          "constraint.")
    return s


def slide_flow(prs):
    s = content_slide(prs, "The flow, with timings",
                      "Package publish to validated production deploy in "
                      "~3 minutes", page=6)

    steps = [
        ("0:00", "Package\npublished", "Artifactory"),
        ("0:02", "Webhook in,\nrule matches", "Event-Driven\nAnsible"),
        ("0:05", "Workflow up,\nroute decided", "Automation\nOrchestrator"),
        ("0:11", "12 of 537\napps affected", "IBM Concert"),
        ("0:24", "Version bumped,\nbranch pushed", "AAP"),
        ("1:40", "Test, build,\npush image", "Tekton"),
        ("2:30", "Auto-sync,\nrolling update", "ArgoCD"),
        ("2:58", "Healthy,\nSBOM updated", "AAP + Concert"),
    ]

    col_w = Inches(1.36)
    gap = Inches(0.20)
    span = len(steps) * col_w + (len(steps) - 1) * gap
    x0 = int((SLIDE_W - span) / 2)
    box_y = Inches(2.68)
    box_h = Inches(1.26)
    rail_y = Inches(2.26)

    rail = slide_rule(s, x0, int(rail_y), int(span))

    for i, (t, what, who) in enumerate(steps):
        x = int(x0 + i * (col_w + gap))
        textbox(s, x, Inches(1.74), col_w, Inches(0.3), t, size=12.5,
                bold=True, color=RED, align=PP_ALIGN.CENTER, font=DISPLAY)
        numbered_circle(s, x + col_w / 2, rail_y, i + 1)
        step_card(s, x, int(box_y), col_w, box_h, what, who, title_size=9.5,
                  body_size=8.5,
                  outline=PLUM if "Concert" in who else RED)

    for i in range(len(steps) - 1):
        xa = int(x0 + i * (col_w + gap)) + col_w
        xb = int(x0 + (i + 1) * (col_w + gap))
        y = int(box_y + box_h / 2)
        arrow(s, xa + Inches(0.02), y, xb - Inches(0.02), y)

    banner = card(s, x0, Inches(4.56), int(span), Inches(0.62), outline=None,
                  fill=RED, radius=0.22)
    tf = banner.text_frame
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = ("~3 minutes  ·  package publish to validated production "
              "deployment  ·  zero manual steps")
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = WHITE
    r.font.name = DISPLAY

    textbox(s, x0, Inches(5.5), int(span), Inches(0.9),
            "Compare: the same change through a ticket-driven process is "
            "typically 2–6 weeks of elapsed time, most of it spent "
            "identifying which applications are affected and who owns them.",
            size=13, color=MUTED, italic=True, align=PP_ALIGN.CENTER,
            spacing=1.2)

    notes(s, "Timings are representative of the lab environment. The point is "
             "the shape of the curve, not the exact seconds.")
    return s


def slide_rule(slide, x, y, w):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, Pt(1.25))
    r.fill.solid()
    r.fill.fore_color.rgb = HAIRLINE
    r.line.fill.background()
    _no_shadow(r)
    r.text_frame.text = ""
    return r


def slide_modules(prs):
    s = content_slide(prs, "What you'll do", "Three modules, 55 minutes "
                      "hands-on", page=7)

    mods = [
        ("01", "DETECT", "~15 min",
         "Publish the Lightwell package. Watch EDA fire. Read Concert's "
         "portfolio impact analysis and Arena View topology, and follow the "
         "AO routing decision."),
        ("02", "ORCHESTRATE", "~25 min",
         "Watch the AO workflow drive the git commit, the Tekton PipelineRun "
         "and the ArgoCD sync. See new pods come up on the remediated "
         "package."),
        ("03", "VALIDATE", "~15 min",
         "Walk the full audit trail — Artifactory › EDA › AO › Concert › "
         "Tekton › ArgoCD › running pod — and confirm Concert's inventory "
         "reflects reality."),
    ]

    x = Inches(1.15)
    w = Inches(3.62)
    for num, name, dur, desc in mods:
        card(s, x, Inches(1.95), w, Inches(3.55), outline=HAIRLINE, fill=WHITE)
        textbox(s, x + Inches(0.32), Inches(2.22), Inches(1.5), Inches(0.75),
                num, size=38, bold=True, color=RGBColor(0xEB, 0xEB, 0xEB),
                font=DISPLAY)
        textbox(s, x + Inches(0.32), Inches(3.02), w - Inches(0.64),
                Inches(0.35), name, size=18, bold=True, color=RED,
                font=DISPLAY)
        textbox(s, x + Inches(0.32), Inches(3.42), w - Inches(0.64),
                Inches(0.28), dur, size=11.5, bold=True, color=MUTED)
        textbox(s, x + Inches(0.32), Inches(3.84), w - Inches(0.64),
                Inches(1.4), desc, size=12, spacing=1.2)
        x += w + Inches(0.24)

    notes(s, "Module 1 is the interesting one conceptually; module 2 is where "
             "it gets satisfying to watch.")
    return s


def slide_watch(prs):
    s = content_slide(prs, "What to watch for",
                      "Four moments that carry the whole argument", page=8)

    items = [
        ("The switch node",
         "That is where your policy lives. Change one condition and the whole "
         "risk posture changes."),
        ("The Concert query latency",
         "One API call replaces a portfolio-wide scan. That is the entire "
         "value proposition in a single HTTP request."),
        ("The ArgoCD sync",
         "Nobody SSH'd anywhere. Desired state went into Git and the cluster "
         "converged. Rollback is one command."),
        ("The audit trail",
         "Every hop is independently queryable. This is what an auditor "
         "actually asks for."),
    ]

    y = Inches(1.82)
    for head, body in items:
        slide_chevron(s, Inches(1.15), y, Inches(4.55), Inches(0.92), head,
                      size=14)
        textbox(s, Inches(6.05), y + Inches(0.10), Inches(6.2), Inches(0.75),
                body, size=13, color=RED, spacing=1.15)
        y += Inches(1.18)

    notes(s, "Call these out live as they happen rather than reading the "
             "slide up front.")
    return s


def slide_questions(prs):
    s = content_slide(prs, "Demo questions",
                      "Where should this go next? We want your opinions.",
                      page=9)

    qs = [
        ("Wrap it all in a ServiceNow ticket?",
         "Auto-create a change record at the switch node, attach the Concert "
         "impact analysis and Tekton logs, close it on validation. "
         "Governance, or just latency?"),
        ("More OpenShift-native — or more pizzazz?",
         "Should more of this run as OCP-native primitives? What would make "
         "the demo memorable rather than merely correct?"),
        ("Policy gates.",
         "Where do OPA / Kyverno / ACS admission policies belong? Block at "
         "admission, or gate earlier in the pipeline? Who owns the policy?"),
        ("Red Hat Dependency Analytics.",
         "Does RHDA overlap with Concert here, or complement it? Shift-left "
         "in the IDE vs. portfolio-wide inventory — do both run?"),
        ("Trusted Profile Analyzer.",
         "Where does TPA fit for SBOM attestation and VEX? Should Tekton emit "
         "a signed SBOM so the trail includes provenance, not just a version "
         "bump?"),
        ("Should prod ever be fully hands-off?",
         "The one we keep arguing about. If an approval gate is always "
         "required — approval by whom, and on what evidence?"),
    ]

    colx = [Inches(1.15), Inches(7.05)]
    colw = Inches(5.15)
    for i, (head, body) in enumerate(qs):
        x = colx[i % 2]
        y = Inches(1.78) + Inches(1.62) * (i // 2)
        numbered_circle(s, x + Inches(0.17), y + Inches(0.18), i + 1,
                        d=Inches(0.30))
        textbox(s, x + Inches(0.46), y, colw - Inches(0.46), Inches(0.3),
                head, size=14, bold=True, color=RED, font=DISPLAY)
        textbox(s, x + Inches(0.46), y + Inches(0.36), colw - Inches(0.46),
                Inches(1.05), body, size=11.5, color=INK, spacing=1.12)

    notes(s, "Leave real time for this. The ServiceNow and policy-gate "
             "questions usually generate the most disagreement.")
    return s


def closing(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    if os.path.exists(TITLE_BG):
        s.shapes.add_picture(TITLE_BG, 0, 0, SLIDE_W, SLIDE_H)
    textbox(s, Inches(0.78), Inches(2.55), Inches(9.6), Inches(1.0),
            "Let's build it", size=44, color=WHITE, font=DISPLAY)
    textbox(s, Inches(0.80), Inches(3.75), Inches(9.6), Inches(0.4),
            "Module 1: Detect — package publish and Concert impact analysis",
            size=17, color=WHITE, font=TEXT)
    notes(s, "Hand off to the lab guide and start Module 1.")
    return s


def main():
    out = sys.argv[1] if len(sys.argv) > 1 else "Lab-Intro-Presentation.pptx"
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    cover(prs)
    slide_problem(prs)
    slide_inversion(prs)
    slide_cast(prs)
    slide_architecture(prs)
    slide_flow(prs)
    slide_modules(prs)
    slide_watch(prs)
    slide_questions(prs)
    closing(prs)

    prs.save(out)
    print(f"Wrote {out} ({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    main()
